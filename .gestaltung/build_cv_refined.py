from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\Jordi Buscaglia\Documents\Bewerbung')
OUT=ROOT/'Lebenslauf_Jordi_Buscaglia_Formatiert.docx'
INK='000000'; ORANGE='FF5C34'; BLUE='D7EFFF'
doc=Document(); sec=doc.sections[0]
sec.page_width=Cm(21);sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.15);sec.bottom_margin=Cm(1.1)
sec.left_margin=Cm(1.05);sec.right_margin=Cm(1.2)
sec.header_distance=Cm(.3);sec.footer_distance=Cm(.3)
for key in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[key];st.font.name='Calibri';st.font.color.rgb=RGBColor.from_string(INK)
    st.paragraph_format.space_before=Pt(0);st.paragraph_format.space_after=Pt(0)
    st.paragraph_format.line_spacing=1.04
    st.paragraph_format.widow_control=True
for b in list(doc.styles.element.iter(qn('w:pBdr'))): b.getparent().remove(b)
doc.styles['Normal'].font.size=Pt(10.2)
doc.styles['Title'].font.size=Pt(28);doc.styles['Title'].font.bold=True
doc.styles['Heading 1'].font.size=Pt(11.8);doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(8)
doc.styles['Heading 2'].font.size=Pt(10.8);doc.styles['Heading 2'].font.bold=True
doc.styles['Subtitle'].font.italic=False

cols=sec._sectPr.find(qn('w:cols'))
cols.set(qn('w:num'),'2');cols.set(qn('w:equalWidth'),'0');cols.set(qn('w:space'),'595')
for width,space in [(2721,595),(7313,0)]:
    col=OxmlElement('w:col');col.set(qn('w:w'),str(width));col.set(qn('w:space'),str(space));cols.append(col)

def rectangle(name,y,w,h,color):
    sec.header.paragraphs[0].add_run()._r.append(parse_xml(f'''<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"><v:rect id="{name}" o:allowincell="f" fillcolor="#{color}" stroked="f" style="position:absolute;margin-left:0cm;margin-top:{y}cm;width:{w}cm;height:{h}cm;z-index:-251654144;mso-position-horizontal-relative:page;mso-position-vertical-relative:page"><v:fill color="#{color}"/><v:stroke on="f"/></v:rect></w:pict>'''))
rectangle('Sidebar',0,6.35,29.7,BLUE)
rectangle('TopAccent',0,21,.17,ORANGE)

def p(text='',size=None,bold=False,after=0,before=0,keep=False):
    out=doc.add_paragraph();f=out.paragraph_format
    f.space_after=Pt(after);f.space_before=Pt(before);f.keep_with_next=keep
    r=out.add_run(text);r.bold=bold
    if size:r.font.size=Pt(size)
    return out
def heading(text,before=15):
    out=doc.add_paragraph(text,style='Heading 1');out.paragraph_format.space_before=Pt(before)
    return out

heading('KONTAKT',before=0)
p('Am Elisabethgehölz 6\n20535 Hamburg',size=10,after=7)
p('+49 157 53646577',size=10,after=7)
p('Jordi.buscaglia577@gmail.com',size=9.1,after=4)
heading('KENNTNISSE',before=19)
for label,text in [('Programmierung','Python, C'),('Software','MATLAB, Simulink,\nAltium, LTspice'),('Weitere Werkzeuge','Git, GitHub, Codex')]:
    p(label,bold=True,size=10,after=3,keep=True)
    p(text,size=10,after=11)
heading('SPRACHEN',before=10)
for label,text in [('Deutsch','Muttersprache'),('Englisch','Verhandlungssicher'),('Spanisch','Grundkenntnisse')]:
    p(label,bold=True,size=10,after=3,keep=True)
    p(text,size=10,after=11)

br=p(size=1);br.paragraph_format.line_spacing=Pt(1);br.add_run().add_break(WD_BREAK.COLUMN)
title=doc.add_paragraph('Jordi Buscaglia',style='Title');title.paragraph_format.space_after=Pt(16)
heading('BERUFSERFAHRUNG',before=0)

RAG='Forschung zu Retrieval-Augmented Generation (RAG) mit Fokus auf die Verbesserung des Tabellen-Parsings und die Implementierung neuer Verfahren.'
LLM='Machbarkeitsstudien zur Codegenerierung mit lokalen Sprachmodellen wie Qwen, Llama und Mistral.'

def employer(name,city,before=0):
    out=p('',before=before,after=7,keep=True)
    out.paragraph_format.tab_stops.add_tab_stop(Cm(12.9),WD_TAB_ALIGNMENT.RIGHT)
    out.add_run(name).bold=True
    r=out.add_run('\t'+city);r.font.size=Pt(9.5)
def role(date,title,items,before=0):
    p(title,bold=True,size=10.5,before=before,after=3,keep=True)
    p(date,size=9.2,after=5,keep=True)
    for item in items:
        out=p('',after=3)
        out.paragraph_format.left_indent=Cm(.28)
        out.paragraph_format.first_line_indent=Cm(-.28)
        out.add_run('• ').font.color.rgb=RGBColor.from_string(ORANGE)
        out.add_run(item)

employer('Fraunhofer-Institut für Siliziumtechnologie','Itzehoe')
role('Jun. 2026 - heute','Masterand im Bereich Mikroelektronische Systeme',[RAG,LLM])
role('Sep. 2025 - Jun. 2026','Werkstudent im Bereich optische Systeme',[RAG,LLM],before=7)
employer('Philips Healthcare','Hamburg',before=11)
role('Jan. 2025 - Sep. 2025','Bachelorand im Bereich Advanced Research and Development',[
    'Bachelorarbeit: „Large Language Models für die Softwareentwicklung: Evaluierung der Machbarkeit lokaler Large Language Models zur Codegenerierung“.'])
role('Okt. 2024 - Jan. 2025','Praktikant im Bereich Advanced Research and Development',[RAG,LLM],before=7)
role('Jan. 2024 - Sep. 2024','Werkstudent in der Testautomatisierung von Röntgendetektoren',[
    'Erstellung und Durchführung von Testplänen und Dauerlauftests; Anpassung der Teststände, Protokollierung von Messdaten und Nachbearbeitung von Röntgenaufnahmen zur Qualitätssicherung.',
    'Erfahrung im technischen Projektmanagement und in der Zusammenarbeit im internationalen Konzernumfeld.'],before=7)

heading('BILDUNGSWEG',before=13)
def education(date,title,detail):
    out=p('',after=2,keep=True)
    out.paragraph_format.tab_stops.add_tab_stop(Cm(12.9),WD_TAB_ALIGNMENT.RIGHT)
    out.add_run(title).bold=True
    r=out.add_run('\t'+date);r.font.size=Pt(9.2)
    p(detail,size=10,after=7)
education('2025 - heute','Mikroelektronische Systeme (M.Sc.)','Hochschule für Angewandte Wissenschaften, Hamburg')
education('2020 - 2025','Mechatronik (B.Sc.)','Hochschule für Angewandte Wissenschaften, Hamburg\nSchwerpunkt Elektrische Antriebe und Robotik')
education('2011 - 2019','Margaretha-Rothe-Gymnasium','Hamburg\nLeistungskurse: Mathematik, Physik, Biologie und Englisch')

doc.core_properties.title='Lebenslauf Jordi Buscaglia'
doc.core_properties.subject='Lebenslauf';doc.core_properties.author='Jordi Buscaglia'
doc.core_properties.comments='';doc.core_properties.keywords=''
doc.save(OUT)
print(OUT)
