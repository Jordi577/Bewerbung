from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

ROOT=Path(r'C:\Users\Jordi Buscaglia\Documents\Bewerbung')
OUT=ROOT/'Lebenslauf_Jordi_Buscaglia_Design_2.docx'
INK='351E28'; ORANGE='FF5C34'; BLUE='D7EFFF'
doc=Document(); s=doc.sections[0]
s.page_width=Cm(21); s.page_height=Cm(29.7)
s.top_margin=Cm(1.15); s.bottom_margin=Cm(1.15)
s.left_margin=Cm(1.5); s.right_margin=Cm(1.5)
s.header_distance=Cm(.3); s.footer_distance=Cm(.3)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
    st=doc.styles[name]; st.font.name='Calibri'; st.font.color.rgb=RGBColor.from_string(INK)
    st.paragraph_format.space_before=Pt(0); st.paragraph_format.space_after=Pt(0)
    st.paragraph_format.line_spacing=1.04
for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.2)
doc.styles['Title'].font.size=Pt(33); doc.styles['Title'].font.bold=True
doc.styles['Title'].font.color.rgb=RGBColor.from_string(BLUE)
doc.styles['Heading 1'].font.size=Pt(12.2); doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 1'].paragraph_format.space_before=Pt(11)
doc.styles['Heading 1'].paragraph_format.space_after=Pt(6)
doc.styles['Heading 2'].font.size=Pt(10.5); doc.styles['Heading 2'].font.bold=True

def background(ident,x,y,w,h,color):
    s.header.paragraphs[0].add_run()._r.append(parse_xml(f'''<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"><v:rect id="{ident}" o:allowincell="f" fillcolor="#{color}" stroked="f" style="position:absolute;margin-left:{x}cm;margin-top:{y}cm;width:{w}cm;height:{h}cm;z-index:-251654144;mso-position-horizontal-relative:page;mso-position-vertical-relative:page"><v:fill color="#{color}"/><v:stroke on="f"/></v:rect></w:pict>'''))
background('Header',0,0,21,5.5,INK)
background('Accent',0,1.15,.35,2.8,ORANGE)

p=doc.add_paragraph(style='Title')
shape=p.add_run().add_picture(str(ROOT/'foto.png'),width=Cm(4),height=Cm(4))
inline=shape._inline; bf=inline.find('.//'+qn('pic:blipFill'))
crop=OxmlElement('a:srcRect');crop.set('l','9375');crop.set('r','9375');bf.insert(1,crop)
anchor=parse_xml('''<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="1" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV></wp:anchor>''')
for ch in list(inline):anchor.append(ch)
anchor.insert(4,OxmlElement('wp:wrapNone'));inline.getparent().replace(inline,anchor)
anchor.find(qn('wp:docPr')).set('descr','Bewerbungsfoto Jordi Buscaglia')
p.add_run('Jordi Buscaglia');p.paragraph_format.space_after=Pt(4)
p=doc.add_paragraph();r=p.add_run('LEBENSLAUF');r.font.size=Pt(10);r.font.color.rgb=RGBColor.from_string(BLUE)
p.paragraph_format.space_after=Pt(13)
for text in ['Am Elisabethgehölz 6, 20535 Hamburg','+49 157 53646577','Jordi.buscaglia577@gmail.com']:
    p=doc.add_paragraph();r=p.add_run(text);r.font.color.rgb=RGBColor.from_string(BLUE);r.font.size=Pt(10.5)
    p.paragraph_format.space_after=Pt(2)
p=doc.add_paragraph();p.paragraph_format.line_spacing=Pt(1);p.paragraph_format.space_after=Pt(32)

def heading(text,first=False):
    p=doc.add_paragraph(text,style='Heading 1')
    if first:p.paragraph_format.space_before=Pt(0)
    return p
def para(text,after=0,indent=3.2,bold=False,keep=False):
    p=doc.add_paragraph();p.paragraph_format.left_indent=Cm(indent)
    p.paragraph_format.space_after=Pt(after);p.paragraph_format.keep_with_next=keep
    p.paragraph_format.widow_control=True;p.add_run(text).bold=bold
    return p
def employer(text,city,before=0):
    p=para('',after=4,bold=True,keep=True);p.paragraph_format.space_before=Pt(before)
    p.add_run(text).bold=True;p.add_run('  |  '+city)
def dated(date,title,before=0,after=4):
    p=para('',after=after,keep=True);p.paragraph_format.space_before=Pt(before)
    p.paragraph_format.first_line_indent=Cm(-3.2)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(3.2),WD_TAB_ALIGNMENT.LEFT)
    r=p.add_run(date);r.font.size=Pt(9.2)
    p.add_run('\t');p.add_run('● ').font.color.rgb=RGBColor.from_string(ORANGE)
    p.add_run(title).bold=True
    return p
def bullet(text):
    p=para('',after=3,indent=3.45)
    p.paragraph_format.first_line_indent=Cm(-.25)
    p.add_run('• ').font.color.rgb=RGBColor.from_string(ORANGE);p.add_run(text)
def role(date,title,items,before=0):
    dated(date,title,before)
    for item in items:bullet(item)

RAG='Forschung zu Retrieval-Augmented Generation (RAG) mit Fokus auf die Verbesserung des Tabellen-Parsings und die Implementierung neuer Verfahren.'
LLM='Machbarkeitsstudien zur Codegenerierung mit lokalen Sprachmodellen wie Qwen, Llama und Mistral.'
heading('Berufserfahrung',True)
employer('Fraunhofer-Institut für Siliziumtechnologie','Itzehoe')
role('Jun. 2026 - heute','Masterand im Bereich Mikroelektronische Systeme',[RAG,LLM])
role('Sep. 2025 - Jun. 2026','Werkstudent im Bereich optische Systeme',[RAG,LLM],before=5)
employer('Philips Healthcare','Hamburg',before=7)
role('Jan. 2025 - Sep. 2025','Bachelorand im Bereich Advanced Research and Development',[
    'Bachelorarbeit: „Large Language Models für die Softwareentwicklung: Evaluierung der Machbarkeit lokaler Large Language Models zur Codegenerierung“.'])
role('Okt. 2024 - Jan. 2025','Praktikant im Bereich Advanced Research and Development',[RAG,LLM],before=5)
role('Jan. 2024 - Sep. 2024','Werkstudent in der Testautomatisierung von Röntgendetektoren',[
    'Erstellung und Durchführung von Testplänen und Dauerlauftests; Anpassung der Teststände, Protokollierung von Messdaten und Nachbearbeitung von Röntgenaufnahmen zur Qualitätssicherung.',
    'Erfahrung im technischen Projektmanagement und in der Zusammenarbeit im internationalen Konzernumfeld.'],before=5)

heading('Bildungsweg')
dated('2025 - heute','Mikroelektronische Systeme (M.Sc.)',after=2)
para('Hochschule für Angewandte Wissenschaften, Hamburg',after=5)
dated('2020 - 2025','Mechatronik (B.Sc.)',after=2)
para('Hochschule für Angewandte Wissenschaften, Hamburg',after=1,keep=True)
para('Schwerpunkt Elektrische Antriebe und Robotik',after=5)
dated('2011 - 2019','Margaretha-Rothe-Gymnasium, Hamburg',after=2)
para('Leistungskurse: Mathematik, Physik, Biologie und Englisch')

heading('Kenntnisse und Sprachen')
for label,text in [
    ('Programmierung','Python, C'),('Software','MATLAB, Simulink, Altium, LTspice'),
    ('Weitere Werkzeuge','Git, GitHub, Codex'),
    ('Sprachen','Deutsch: Muttersprache  |  Englisch: Verhandlungssicher  |  Spanisch: Grundkenntnisse')]:
    p=para('',indent=0,after=3);p.paragraph_format.tab_stops.add_tab_stop(Cm(3.2))
    p.add_run(label).bold=True;p.add_run('\t'+text)

doc.core_properties.title='Lebenslauf Jordi Buscaglia'
doc.core_properties.subject='Lebenslauf';doc.core_properties.author='Jordi Buscaglia'
doc.core_properties.comments='';doc.save(OUT)
print(OUT)
