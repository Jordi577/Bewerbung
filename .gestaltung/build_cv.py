from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn

ROOT = Path(r'C:\Users\Jordi Buscaglia\Documents\Bewerbung')
OUT = ROOT / 'Lebenslauf_Jordi_Buscaglia.docx'
INK = '351E28'
ACCENT = 'FF5C34'
PALE = 'D7EFFF'
doc = Document()
s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.top_margin, s.bottom_margin = Cm(1.5), Cm(1.4)
s.left_margin, s.right_margin = Cm(1.5), Cm(1.5)
s.header_distance, s.footer_distance = Cm(.5), Cm(.5)

for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet']:
    st = doc.styles[sty]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor.from_string(INK)
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.space_before = Pt(0)
    st.paragraph_format.line_spacing = 1.05
doc.styles['Normal'].font.size = Pt(10.3)
doc.styles['Title'].font.size = Pt(29)
doc.styles['Title'].font.bold = True
doc.styles['Subtitle'].font.size = Pt(10)
doc.styles['Heading 1'].font.size = Pt(12)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(8)
doc.styles['Heading 2'].font.size = Pt(11)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
doc.styles['Subtitle'].font.italic = False
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)

# Background decoration lives in the header and does not interrupt editing
# or the reading order of the CV text.
header_run = s.header.paragraphs[0].add_run()
for ident, top, width, height, fill in [
    ('PaletteTop', '0cm', '21cm', '0.18cm', ACCENT),
    ('PaletteSidebar', '6.25cm', '6.55cm', '23.45cm', PALE),
]:
    header_run._r.append(parse_xml(f'''<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"><v:rect id="{ident}" o:allowincell="f" fillcolor="#{fill}" stroked="f" style="position:absolute;margin-left:0cm;margin-top:{top};width:{width};height:{height};z-index:-251654144;mso-position-horizontal-relative:page;mso-position-vertical-relative:page"><v:fill color="#{fill}"/><v:stroke on="f"/></v:rect></w:pict>'''))

# Keep the photo replaceable as a normal Word picture. Its crop is stored
# in DrawingML; the original bitmap remains embedded without modification.
p = doc.add_paragraph(style='Title')
r = p.add_run()
shape = r.add_picture(str(ROOT/'foto.png'),width=Cm(3.2),height=Cm(4.0))
inline = shape._inline
blipfill = inline.find('.//'+qn('pic:blipFill'))
crop = OxmlElement('a:srcRect')
crop.set('l','17500'); crop.set('r','17500')
blipfill.insert(1,crop)
anchor = parse_xml('''<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="0" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="margin"><wp:align>right</wp:align></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV></wp:anchor>''')
for child in list(inline): anchor.append(child)
anchor.insert(4,OxmlElement('wp:wrapNone'))
inline.getparent().replace(inline,anchor)
anchor.find(qn('wp:docPr')).set('descr','Bewerbungsfoto Jordi Buscaglia')
p.add_run('Jordi Buscaglia')
p.paragraph_format.space_after = Pt(5)
p = doc.add_paragraph('LEBENSLAUF',style='Subtitle')
p.paragraph_format.space_after = Pt(14)
for txt in ['Am Elisabethgehölz 6, 20535 Hamburg', '+49 157 53646577', 'Jordi.buscaglia577@gmail.com']:
    p=doc.add_paragraph(txt)
    p.paragraph_format.space_after=Pt(3)
    for run in p.runs: run.font.size=Pt(10.5)

# Fixed whitespace for a calm header, followed by editable newspaper columns.
p=doc.add_paragraph()
p.paragraph_format.line_spacing=Pt(1)
p.paragraph_format.space_after=Pt(27)
p.paragraph_format.space_before=Pt(0)
sec=doc.add_section(WD_SECTION_START.CONTINUOUS)
cols=sec._sectPr.find(qn('w:cols'))
cols.set(qn('w:num'),'2'); cols.set(qn('w:equalWidth'),'0')
cols.set(qn('w:space'),'454')
for w,sp in [(2721,454),(7030,0)]:
    c=OxmlElement('w:col'); c.set(qn('w:w'),str(w)); c.set(qn('w:space'),str(sp)); cols.append(c)

def paragraph(text='',bold=False,size=None,color=None,after=0,before=0,keep=False):
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(after)
    p.paragraph_format.space_before=Pt(before)
    p.paragraph_format.keep_with_next=keep
    p.paragraph_format.widow_control=True
    r=p.add_run(text); r.bold=bold
    if size:r.font.size=Pt(size)
    if color:r.font.color.rgb=RGBColor.from_string(color)
    return p

def heading(text,first=False):
    p=doc.add_paragraph(text,style='Heading 1')
    if first:p.paragraph_format.space_before=Pt(0)
    return p

heading('Bildungsweg',True)
def education(date,degree,institution,detail=None):
    paragraph(date,size=9.4,color=INK,after=3,keep=True)
    paragraph(degree,bold=True,after=3,keep=True)
    paragraph(institution,after=3 if detail else 11,keep=bool(detail))
    if detail: paragraph(detail,after=11)

education('2025 - heute','Mikroelektronische Systeme (M.Sc.)','Hochschule für Angewandte Wissenschaften, Hamburg')
education('2020 - 2025','Mechatronik (B.Sc.)','Hochschule für Angewandte Wissenschaften, Hamburg','Schwerpunkt Elektrische Antriebe und Robotik')
education('2011 - 2019','Margaretha-Rothe-Gymnasium','Hamburg','Leistungskurse: Mathematik, Physik, Biologie und Englisch')

heading('Kenntnisse')
for lab,txt in [('Programmierung','Python, C'),('Software','MATLAB, Simulink, Altium, LTspice'),('Weitere Werkzeuge','Git, GitHub, Codex')]:
    paragraph(lab,bold=True,after=3,keep=True)
    paragraph(txt,after=10)
heading('Sprachen')
for lab,txt in [('Deutsch','Muttersprache'),('Englisch','Verhandlungssicher'),('Spanisch','Grundkenntnisse')]:
    paragraph(lab,bold=True,after=2,keep=True)
    paragraph(txt,after=8)

p=doc.add_paragraph(); p.paragraph_format.line_spacing=Pt(1)
p.add_run().add_break(WD_BREAK.COLUMN)
heading('Berufserfahrung',True)

RAG='Forschung zu Retrieval-Augmented Generation (RAG) mit Fokus auf die Verbesserung des Tabellen-Parsings und die Implementierung neuer Verfahren.'
LLM='Machbarkeitsstudien zur Codegenerierung mit lokalen Sprachmodellen wie Qwen, Llama und Mistral.'
def employer(name,city,before=0):
    p=doc.add_paragraph(style='Heading 2')
    p.paragraph_format.space_before=Pt(before)
    p.add_run(name)
    p.add_run(' | '+city).bold=False

def role(date,title,items,before=0):
    paragraph(date,size=9.4,color=INK,before=before,after=3,keep=True)
    paragraph(title,bold=True,after=5,keep=True)
    for i,item in enumerate(items):
        p=paragraph(item,after=3)
        p.paragraph_format.left_indent=Cm(.25)
        p.paragraph_format.first_line_indent=Cm(-.25)
        p.runs[0].text=''
        bullet=p.add_run('• ')
        bullet.font.color.rgb=RGBColor.from_string(ACCENT)
        p.add_run(item)

employer('Fraunhofer-Institut für Siliziumtechnologie','Itzehoe')
role('Jun. 2026 - heute','Masterand im Bereich Mikroelektronische Systeme',[RAG,LLM])
role('Sep. 2025 - Jun. 2026','Werkstudent im Bereich optische Systeme',[RAG,LLM],before=7)
employer('Philips Healthcare','Hamburg',before=12)
role('Jan. 2025 - Sep. 2025','Bachelorand im Bereich Advanced Research and Development',['Bachelorarbeit: „Large Language Models für die Softwareentwicklung: Evaluierung der Machbarkeit lokaler Large Language Models zur Codegenerierung“.'])
role('Okt. 2024 - Jan. 2025','Praktikant im Bereich Advanced Research and Development',[RAG,LLM],before=7)
role('Jan. 2024 - Sep. 2024','Werkstudent in der Testautomatisierung von Röntgendetektoren',[
    'Erstellung und Durchführung von Testplänen und Dauerlauftests; Anpassung der Teststände, Protokollierung von Messdaten und Nachbearbeitung von Röntgenaufnahmen zur Qualitätssicherung.',
    'Erfahrung im technischen Projektmanagement und in der Zusammenarbeit im internationalen Konzernumfeld.'
],before=7)

doc.core_properties.title='Lebenslauf Jordi Buscaglia'
doc.core_properties.subject='Lebenslauf'
doc.core_properties.author='Jordi Buscaglia'
doc.core_properties.keywords=''
doc.core_properties.comments=''
doc.save(OUT)
print(OUT)

